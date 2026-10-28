"""
Tests del Excel final: genera el reporte en una carpeta temporal, lo vuelve a
abrir con openpyxl y revisa estructura, cifras y estilo.
"""
import math
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from generar_reporte import AVISO, FILAS_GRAFICO, FORMATO_PESOS, HOJAS, SIN_DATO, generar

CARPETA = Path(__file__).resolve().parent.parent / "datos_semana"


@pytest.fixture(scope="module")
def reporte(tmp_path_factory):
    ruta = tmp_path_factory.mktemp("salida") / "reporte.xlsx"
    info = generar(CARPETA, ruta)
    return load_workbook(ruta), info


def buscar(hoja, texto, columna=None):
    """Número de fila de la primera celda con ese texto (en una columna, si se indica)."""
    for fila in hoja.iter_rows():
        for celda in fila:
            if celda.value == texto and (columna is None or celda.column == columna):
                return celda.row
    raise AssertionError(f"No se encontró '{texto}' en la hoja {hoja.title}")


def tabla(hoja, fila_encabezado):
    """Filas de una tabla como diccionarios {encabezado: valor}, hasta la primera fila vacía."""
    encabezados = [c.value for c in hoja[fila_encabezado] if c.value is not None]
    filas = []
    for fila in hoja.iter_rows(min_row=fila_encabezado + 1, max_col=len(encabezados), values_only=True):
        if all(v is None for v in fila):
            break
        filas.append(dict(zip(encabezados, fila)))
    return filas


# --- Estructura ---------------------------------------------------------------

def test_tiene_las_cinco_hojas_en_orden(reporte):
    libro, _ = reporte
    assert libro.sheetnames == HOJAS


def test_cada_hoja_avisa_que_los_datos_son_ficticios(reporte):
    libro, _ = reporte
    for hoja in libro:
        assert hoja["A2"].value == AVISO, hoja.title


def test_ninguna_hoja_muestra_cuadricula(reporte):
    libro, _ = reporte
    assert all(not hoja.sheet_view.showGridLines for hoja in libro)


def test_encabezados_azul_oscuro_con_texto_blanco(reporte):
    libro, _ = reporte
    hoja = libro["Datos limpios"]
    celda = hoja["A4"]
    assert celda.value == "Fecha"
    assert celda.fill.fgColor.rgb.endswith("1F3864")
    assert celda.font.color.rgb.endswith("FFFFFF")


# --- Cifras -------------------------------------------------------------------

def test_total_del_resumen_es_la_suma_de_datos_limpios(reporte):
    libro, info = reporte
    resumen = libro["Resumen"]
    total = resumen.cell(buscar(resumen, "Total cadena", columna=1), 2).value

    datos = tabla(libro["Datos limpios"], 4)
    suma_semana_actual = sum(d["Monto"] for d in datos if d["Semana (lunes)"].date() == info["actual"])
    assert total == suma_semana_actual == info["ventas"] == 15637000


def test_sucursales_suman_el_total(reporte):
    libro, _ = reporte
    filas = tabla(libro["Resumen"], 10)
    sucursales = [f for f in filas if f["Sucursal"] != "Total cadena"]
    total = next(f for f in filas if f["Sucursal"] == "Total cadena")
    assert [f["Sucursal"] for f in sucursales] == ["Providencia", "Ñuñoa", "Maipú"]
    assert sum(f["Semana actual"] for f in sucursales) == total["Semana actual"]
    assert sum(f["Semana anterior"] for f in sucursales) == total["Semana anterior"]


def test_resumen_tiene_cuatro_indicadores_y_no_ticket_promedio(reporte):
    libro, _ = reporte
    hoja = libro["Resumen"]
    etiquetas = [hoja.cell(5, c).value for c in (1, 3, 5, 7)]
    assert etiquetas == ["Ventas totales", "Venta promedio diaria", "Unidades vendidas", "Producto estrella"]
    textos = " ".join(str(c.value) for fila in hoja.iter_rows() for c in fila if c.value).lower()
    assert "ticket" not in textos


def test_datos_limpios_tiene_filtros_y_origen_de_cada_fila(reporte):
    libro, info = reporte
    hoja = libro["Datos limpios"]
    datos = tabla(hoja, 4)
    assert len(datos) == info["filas"]
    assert hoja.auto_filter.ref == f"A4:I{4 + info['filas']}"
    assert all(d["Archivo de origen"] and d["Fila de origen"] > 1 for d in datos)


def secciones_revision(hoja):
    """{'Errores corregidos': fila del título, 'Conversiones de formato': fila del título}."""
    return {
        celda.value.split(" (")[0]: celda.row
        for celda in hoja["A"]
        if isinstance(celda.value, str) and celda.value.startswith(("Errores corregidos", "Conversiones de formato"))
    }


def test_revision_de_datos_separa_errores_de_conversiones(reporte):
    libro, info = reporte
    hoja = libro["Revisión de datos"]
    secciones = secciones_revision(hoja)
    assert secciones["Errores corregidos"] < secciones["Conversiones de formato"]  # primero los errores
    errores = tabla(hoja, secciones["Errores corregidos"] + 1)
    conversiones = tabla(hoja, secciones["Conversiones de formato"] + 1)

    assert sum(f["N.º de filas"] for f in errores) == info["errores"]
    assert sum(f["N.º de filas"] for f in conversiones) == info["conversiones"]
    assert info["errores"] + info["conversiones"] == info["correcciones"]
    assert hoja["A3"].value.startswith(
        f"Errores reales corregidos: {info['errores']}. Conversiones de formato: {info['conversiones']}.")

    texto_errores = " | ".join(f["Problema"] for f in errores)
    for palabra in ["duplicada", "no coincide", "tipeo", "abreviado", "desconocido", "palabras"]:
        assert palabra in texto_errores, palabra
    for palabra in ["fecha", "'$'", "mayúsculas"]:
        assert palabra not in texto_errores, palabra
    texto_conversiones = " | ".join(f["Problema"] for f in conversiones)
    for palabra in ["fecha en formato", "número de Excel", "'$'", "mayúsculas"]:
        assert palabra in texto_conversiones, palabra


def test_revision_de_datos_tiene_todas_las_correcciones_agrupadas(reporte):
    libro, info = reporte
    hoja = libro["Revisión de datos"]
    secciones = secciones_revision(hoja)
    filas = tabla(hoja, secciones["Errores corregidos"] + 1) + tabla(hoja, secciones["Conversiones de formato"] + 1)
    assert len(filas) == info["grupos"]
    assert sum(f["N.º de filas"] for f in filas) == info["correcciones"]
    duplicada = [f for f in filas if "duplicada" in f["Problema"]]
    assert duplicada == [{
        "Archivo": "providencia_2026-10-19.csv", "N.º de filas": 1, "Filas": "51",
        "Problema": "fila duplicada (igual a la fila 50)", "Acción": "eliminada",
    }]


# --- Estilo -------------------------------------------------------------------

def test_montos_con_formato_de_pesos(reporte):
    libro, _ = reporte
    resumen = libro["Resumen"]
    fila_total = buscar(resumen, "Total cadena", columna=1)
    assert resumen.cell(fila_total, 2).number_format == FORMATO_PESOS
    assert resumen["A6"].number_format == FORMATO_PESOS  # indicador de ventas totales
    assert libro["Datos limpios"]["G5"].number_format == FORMATO_PESOS


def test_variacion_verde_si_sube_y_roja_si_baja(reporte):
    libro, _ = reporte
    hoja = libro["Resumen"]
    providencia = hoja.cell(buscar(hoja, "Providencia", columna=1), 4)
    nunoa = hoja.cell(buscar(hoja, "Ñuñoa", columna=1), 4)
    assert providencia.value > 0 and providencia.font.color.rgb.endswith("006100")
    assert nunoa.value < 0 and nunoa.font.color.rgb.endswith("9C0006")


def test_reposicion_colorea_estados_y_solo_sugiere_si_hay_que_pedir(reporte):
    libro, _ = reporte
    hoja = libro["Reposición de insumos"]
    filas = tabla(hoja, 6)
    estados = [f["Estado"] for f in filas]
    # 18 insumos × 3 sucursales = 54 filas; pocas alertas en toda la cadena.
    assert estados.count("URGENTE") == 2 and estados.count("Pedir") == 3 and estados.count("OK") == 49

    colores = {"URGENTE": "FFC7CE", "Pedir": "FFEB9C", "OK": "C6EFCE"}
    for i, f in enumerate(filas):
        assert hoja.cell(7 + i, 8).fill.fgColor.rgb.endswith(colores[f["Estado"]])
        if f["Estado"] == "OK":
            assert f["Cantidad sugerida a pedir"] == SIN_DATO
        else:
            assert f["Cantidad sugerida a pedir"] > 0


def test_reposicion_es_por_sucursal(reporte):
    libro, _ = reporte
    hoja = libro["Reposición de insumos"]
    filas = tabla(hoja, 6)
    por_sucursal = {s: sum(1 for f in filas if f["Sucursal"] == s) for s in ("Providencia", "Ñuñoa", "Maipú")}
    assert por_sucursal == {"Providencia": 18, "Ñuñoa": 18, "Maipú": 18}

    alertas = [(f["Estado"], f["Sucursal"], f["Insumo"]) for f in filas if f["Estado"] != "OK"]
    assert alertas == [
        ("URGENTE", "Providencia", "Palta"), ("URGENTE", "Maipú", "Limón"),
        ("Pedir", "Providencia", "Huevos"), ("Pedir", "Ñuñoa", "Carne de vacuno"), ("Pedir", "Maipú", "Pan"),
    ]
    assert hoja["A4"].value == ("Alertas por sucursal: Providencia: 1 URGENTE, 1 Pedir · "
                                "Ñuñoa: 1 Pedir · Maipú: 1 URGENTE, 1 Pedir")
    assert hoja.auto_filter.ref == "A6:I60"


def test_reposicion_avisa_productos_sin_receta(reporte):
    libro, _ = reporte
    textos = [c.value for fila in libro["Reposición de insumos"].iter_rows() for c in fila if c.value]
    assert any("Promo almuerzo" in str(t) and "no tiene receta" in str(t) for t in textos)


def texto_mostrado(celda):
    """Aproximación del texto que Excel muestra en la celda."""
    valor, formato = celda.value, celda.number_format
    if isinstance(valor, date):
        return "00/00/0000"
    if isinstance(valor, (int, float)):
        if "$" in formato:
            return f"${valor:,.0f}"
        if "%" in formato:
            return f"+{valor * 100:.1f}%"
        return f"{valor:,.1f}" if "0.0" in formato else f"{valor:,.0f}"
    return str(valor)


def test_ningun_texto_queda_cortado(reporte):
    """Revisión aproximada: el largo del texto (según el tamaño de letra) cabe en su celda.

    Un número que no cabe se ve como ###; un texto que no cabe se corta si la celda
    de al lado tiene datos; un texto con ajuste de línea se corta si la fila es baja.
    """
    libro, _ = reporte
    cortados = []
    for hoja in libro:
        anchos = lambda c: hoja.column_dimensions[get_column_letter(c)].width or 8.43
        espacio = {}
        for rango in hoja.merged_cells.ranges:
            espacio[(rango.min_row, rango.min_col)] = sum(anchos(c) for c in range(rango.min_col, rango.max_col + 1))
            for celda in list(rango.cells)[1:]:
                espacio[celda] = None  # parte de una celda combinada: se revisa la primera
        for fila in hoja.iter_rows():
            for celda in fila:
                clave = (celda.row, celda.column)
                if celda.value is None or (clave in espacio and espacio[clave] is None):
                    continue
                ancho = espacio.get(clave) or anchos(celda.column)
                tamano = (celda.font.size or 11) / 11
                largo = len(texto_mostrado(celda)) * tamano * (1.1 if celda.font.bold else 1)
                if celda.alignment.wrap_text:
                    lineas = math.ceil(largo / (ancho - 1))
                    alto = hoja.row_dimensions[celda.row].height or 15
                    if lineas * 15 * tamano > alto:
                        cortados.append(f"{hoja.title}!{celda.coordinate} ({lineas} líneas, alto {alto})")
                elif largo > ancho - 1:
                    es_numero = isinstance(celda.value, (int, float, date))
                    if es_numero or hoja.cell(celda.row, celda.column + 1).value is not None:
                        cortados.append(f"{hoja.title}!{celda.coordinate} '{texto_mostrado(celda)}'")
    assert cortados == []


# --- Gráficos -----------------------------------------------------------------

def test_graficos_debajo_de_la_tabla_sin_tapar_otras_tablas(reporte):
    libro, _ = reporte
    hoja = libro["Resumen"]
    graficos = hoja._charts
    assert len(graficos) == 2

    fila_total = buscar(hoja, "Total cadena", columna=1)
    fila_diarias = buscar(hoja, "Ventas diarias de la cadena", columna=1)
    for grafico in graficos:
        fila_inicio = grafico.anchor._from.row + 1  # openpyxl cuenta las filas desde 0
        assert fila_inicio > fila_total                     # debajo de la tabla por sucursal
        assert fila_inicio + FILAS_GRAFICO <= fila_diarias  # y encima de la tabla diaria


def test_leyendas_arriba_sin_tapar_los_nombres_del_eje(reporte):
    # Abajo, la leyenda tapaba los nombres de las sucursales del gráfico de barras.
    libro, _ = reporte
    for grafico in libro["Resumen"]._charts:
        assert grafico.legend.position == "t", grafico.tagname
        assert not grafico.legend.overlay, grafico.tagname


def test_grafico_de_lineas_no_suaviza(reporte):
    libro, _ = reporte
    lineas = [g for g in libro["Resumen"]._charts if g.tagname == "lineChart"]
    assert len(lineas) == 1
    # Debe quedar escrito como "no suavizar" (<smooth val="0"/>), no omitido.
    assert [serie.smooth for serie in lineas[0].series] == [False, False]
