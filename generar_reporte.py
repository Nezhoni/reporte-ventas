"""
Genera el reporte semanal de ventas de Cocina del Valle en Excel.

Lee las planillas de datos_semana/ (sin modificarlas), las limpia, calcula el
resumen, el ranking y la reposición, y escribe salida/reporte_semanal.xlsx.

Uso:
    python3 generar_reporte.py
    python3 generar_reporte.py --datos otra_carpeta/ --salida otro_reporte.xlsx
"""
import argparse
import math
from datetime import timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from calculos import (
    cargar_inventario,
    productos_sin_receta,
    ranking_productos,
    reposicion,
    resumen_sucursales,
    semanas,
    variacion,
    ventas_diarias,
    DIAS_A_CUBRIR,
    DIAS_MARGEN,
    DIAS_PEDIR,
    DIAS_URGENTE,
    TOTAL,
)
from limpieza import Registro, cargar_catalogo, cargar_ventas, pesos

HOJAS = ["Resumen", "Ranking de productos", "Reposición de insumos", "Revisión de datos", "Datos limpios"]
TITULO = "Reporte semanal de ventas · Cocina del Valle"
AVISO = "DATOS DE EJEMPLO: todas las cifras de este reporte son ficticias y se generaron para demostración."

AZUL_OSCURO = "1F3864"
AZUL_CLARO = "DEEAF6"
GRIS_CLARO = "F2F2F2"
GRIS_TEXTO = "595959"
GRIS_SERIE = "A6A6A6"
NARANJO = "C55A11"
VERDE = "006100"
ROJO = "9C0006"
COLORES_ESTADO = {  # (fondo, texto)
    "URGENTE": ("FFC7CE", "9C0006"),
    "Pedir": ("FFEB9C", "9C5700"),
    "OK": ("C6EFCE", "006100"),
    "Sin dato de stock": ("D9D9D9", "404040"),
}

FORMATO_PESOS = '"$"#,##0'
FORMATO_VARIACION = '+0.0%;-0.0%;0.0%'
FORMATO_FECHA = "dd/mm/yyyy"
FORMATO_ENTERO = "#,##0"
FORMATO_DECIMAL = "#,##0.0"
SIN_DATO = "—"

ALTO_GRAFICO_CM = 7.5
FILAS_GRAFICO = 16  # filas que ocupa un gráfico de 7,5 cm, con margen


# --- Utilidades de formato ----------------------------------------------------

def relleno(color):
    return PatternFill("solid", fgColor=color)


def vacio(valor):
    return valor is None or valor is pd.NA or (isinstance(valor, float) and math.isnan(valor))


def porcentaje_texto(valor, signo=True):
    """0.063 -> '+6,3 %' (con coma decimal, como se escribe en Chile)."""
    texto = f"{valor:+.1%}" if signo else f"{valor:.1%}"
    return texto.replace(".", ",").replace("%", " %")


def colorear_variacion(celda):
    if not vacio(celda.value) and celda.value != 0:
        celda.font = Font(color=VERDE if celda.value > 0 else ROJO, bold=celda.font.bold)


def fijar_anchos(hoja, anchos):
    for letra, ancho in anchos.items():
        hoja.column_dimensions[letra].width = ancho


def lineas_necesarias(texto, ancho, negrita=False):
    """Cuántas líneas ocupa un texto con ajuste de línea en una columna de ese ancho."""
    largo = len(str(texto)) * (1.1 if negrita else 1)
    return max(1, math.ceil(largo / max(1, ancho - 2)))


def ancho_columna(hoja, columna):
    return hoja.column_dimensions[get_column_letter(columna)].width


def encabezado_hoja(hoja, subtitulo, ultima_columna):
    """Título, aviso de datos ficticios y subtítulo en las filas 1 a 3."""
    hoja["A1"] = TITULO
    hoja["A1"].font = Font(size=16, bold=True, color=AZUL_OSCURO)
    hoja["A2"] = AVISO
    hoja["A2"].font = Font(bold=True, color=NARANJO)
    hoja["A3"] = subtitulo
    hoja["A3"].font = Font(italic=True, color=GRIS_TEXTO)
    for fila in (1, 2, 3):
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima_columna)


def escribir_tabla(hoja, fila, columnas, registros, ajustar_texto=False):
    """Escribe una tabla con encabezado azul oscuro y filas alternadas en gris.

    columnas: lista de (título, formato de número o None). Los anchos de columna
    deben estar fijados antes: el alto de cada fila se calcula con ellos para
    que ningún texto quede cortado.
    ajustar_texto: si es True, los textos largos de las filas pasan a varias líneas.
    Devuelve el número de la última fila escrita.
    """
    anchos = [ancho_columna(hoja, j) for j in range(1, len(columnas) + 1)]
    for j, (titulo, _) in enumerate(columnas, 1):
        celda = hoja.cell(fila, j, titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno(AZUL_OSCURO)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[fila].height = 15 * max(
        lineas_necesarias(t, a, negrita=True) for (t, _), a in zip(columnas, anchos))

    for i, valores in enumerate(registros):
        f = fila + 1 + i
        for j, ((_, formato), valor) in enumerate(zip(columnas, valores), 1):
            celda = hoja.cell(f, j, None if vacio(valor) else valor)
            if formato:
                celda.number_format = formato
            if i % 2 == 1:
                celda.fill = relleno(GRIS_CLARO)
            celda.alignment = Alignment(vertical="top", wrap_text=ajustar_texto)
        if ajustar_texto:
            hoja.row_dimensions[f].height = 15 * max(
                lineas_necesarias(v, a) for v, a in zip(valores, anchos) if not vacio(v))
    return fila + len(registros)


def rango_semana(lunes):
    return f"{lunes:%d/%m/%Y} al {lunes + timedelta(days=6):%d/%m/%Y}"


# --- Hoja Resumen -------------------------------------------------------------

def texto_variacion(valor):
    """(texto, color) para la línea bajo cada indicador."""
    if vacio(valor):
        return "sin semana anterior para comparar", GRIS_TEXTO
    return f"{porcentaje_texto(valor)} vs semana anterior", VERDE if valor >= 0 else ROJO


def escribir_indicadores(hoja, fila, resumen, ranking):
    total = resumen.iloc[-1]
    estrella = ranking.iloc[0]
    indicadores = [
        ("Ventas totales", total["ventas"], FORMATO_PESOS, texto_variacion(total["variacion"])),
        ("Venta promedio diaria", total["promedio_diario"], FORMATO_PESOS,
         texto_variacion(variacion(total["promedio_diario"], total["promedio_diario_anterior"]))),
        ("Unidades vendidas", total["unidades"], FORMATO_ENTERO,
         texto_variacion(variacion(total["unidades"], total["unidades_anterior"]))),
        ("Producto estrella", estrella["producto"], None,
         (f"{pesos(estrella['monto'])} · {porcentaje_texto(estrella['porcentaje'], signo=False)} de las ventas",
          GRIS_TEXTO)),
    ]
    for k, (etiqueta, valor, formato, (detalle, color)) in enumerate(indicadores):
        columna = 1 + 2 * k
        for desplazamiento, contenido, fuente in [
            (0, etiqueta, Font(size=10, color=GRIS_TEXTO)),
            (1, valor, Font(size=18, bold=True, color=AZUL_OSCURO)),
            (2, detalle, Font(size=10, color=color)),
        ]:
            f = fila + desplazamiento
            celda = hoja.cell(f, columna, contenido)
            celda.font = fuente
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if desplazamiento == 1 and formato:
                celda.number_format = formato
            hoja.merge_cells(start_row=f, start_column=columna, end_row=f, end_column=columna + 1)
            for c in (columna, columna + 1):
                hoja.cell(f, c).fill = relleno(AZUL_CLARO)
    hoja.row_dimensions[fila + 1].height = 30
    hoja.row_dimensions[fila + 2].height = 30


def hoja_resumen(hoja, resumen, diarias, ranking, actual, anterior):
    fijar_anchos(hoja, {"A": 22, "B": 17, "C": 17, "D": 13, "E": 13, "F": 12, "G": 17, "H": 20})
    subtitulo = f"Semana del {rango_semana(actual)}"
    if anterior is not None:
        subtitulo += f", comparada con la semana del {rango_semana(anterior)}"
    encabezado_hoja(hoja, subtitulo, 8)
    escribir_indicadores(hoja, 5, resumen, ranking)

    hoja["A9"] = "Ventas por sucursal"
    hoja["A9"].font = Font(size=12, bold=True, color=AZUL_OSCURO)
    fila_tabla = 10
    columnas = [
        ("Sucursal", None), ("Semana actual", FORMATO_PESOS), ("Semana anterior", FORMATO_PESOS),
        ("Variación", FORMATO_VARIACION), ("Unidades", FORMATO_ENTERO), ("Días con ventas", "0"),
        ("Promedio diario", FORMATO_PESOS), ("Producto estrella", None),
    ]
    registros = resumen[["sucursal", "ventas", "ventas_anterior", "variacion", "unidades",
                         "dias_con_ventas", "promedio_diario", "producto_estrella"]].values.tolist()
    ultima = escribir_tabla(hoja, fila_tabla, columnas, registros)
    for c in range(1, len(columnas) + 1):  # fila del total en negrita
        hoja.cell(ultima, c).font = Font(bold=True)
    for f in range(fila_tabla + 1, ultima + 1):
        colorear_variacion(hoja.cell(f, 4))

    # Gráficos debajo de la tabla, lado a lado.
    fila_graficos = ultima + 2
    ultima_sucursal = ultima - 1  # sin la fila del total
    barras = BarChart()
    barras.type = "col"
    barras.title = "Ventas por sucursal"
    barras.y_axis.numFmt = FORMATO_PESOS
    barras.y_axis.majorGridlines = None
    for columna, color in [(3, GRIS_SERIE), (2, AZUL_OSCURO)]:  # anterior primero, actual después
        barras.add_data(Reference(hoja, min_col=columna, min_row=fila_tabla, max_row=ultima_sucursal),
                        titles_from_data=True)
        barras.series[-1].graphicalProperties.solidFill = color
        barras.series[-1].graphicalProperties.line.solidFill = color
    barras.set_categories(Reference(hoja, min_col=1, min_row=fila_tabla + 1, max_row=ultima_sucursal))
    preparar_grafico(barras, ancho_cm=12.5)
    hoja.add_chart(barras, f"A{fila_graficos}")

    # Tabla de ventas diarias debajo de los gráficos (es la fuente del gráfico de líneas).
    fila_diarias = fila_graficos + FILAS_GRAFICO
    hoja.cell(fila_diarias, 1, "Ventas diarias de la cadena").font = Font(size=12, bold=True, color=AZUL_OSCURO)
    encabezado_diarias = fila_diarias + 1
    ultima_diaria = escribir_tabla(
        hoja, encabezado_diarias,
        [("Día", None), ("Fecha", FORMATO_FECHA), ("Semana actual", FORMATO_PESOS), ("Semana anterior", FORMATO_PESOS)],
        diarias[["dia", "fecha", "ventas", "ventas_anterior"]].values.tolist(),
    )

    lineas = LineChart()
    lineas.title = "Ventas diarias de la cadena"
    lineas.y_axis.numFmt = FORMATO_PESOS
    lineas.y_axis.majorGridlines = None
    for columna, color in [(4, GRIS_SERIE), (3, AZUL_OSCURO)]:
        lineas.add_data(Reference(hoja, min_col=columna, min_row=encabezado_diarias, max_row=ultima_diaria),
                        titles_from_data=True)
        serie = lineas.series[-1]
        serie.smooth = False  # líneas rectas entre puntos: muestran los datos reales
        serie.graphicalProperties.line.solidFill = color
        serie.marker.symbol = "circle"
        serie.marker.graphicalProperties.solidFill = color
        serie.marker.graphicalProperties.line.solidFill = color
    lineas.set_categories(Reference(hoja, min_col=1, min_row=encabezado_diarias + 1, max_row=ultima_diaria))
    preparar_grafico(lineas, ancho_cm=11.8)
    hoja.add_chart(lineas, f"E{fila_graficos}")


def preparar_grafico(grafico, ancho_cm):
    grafico.width = ancho_cm
    grafico.height = ALTO_GRAFICO_CM
    grafico.legend.position = "b"
    # Sin esto, Excel puede ocultar los ejes de los gráficos creados con openpyxl.
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False


# --- Hojas de detalle ---------------------------------------------------------

def texto_cambio(posicion, anterior):
    if vacio(anterior):
        return "Nuevo", GRIS_TEXTO
    diferencia = int(anterior) - int(posicion)
    if diferencia > 0:
        return f"▲ {diferencia}", VERDE
    if diferencia < 0:
        return f"▼ {-diferencia}", ROJO
    return "=", GRIS_TEXTO


def hoja_ranking(hoja, ranking, actual, anterior):
    anchos = [10, 22, 18, 11, 14, 11, 16, 10]
    fijar_anchos(hoja, dict(zip("ABCDEFGH", anchos)))
    encabezado_hoja(hoja, f"Productos ordenados por ventas de la semana del {rango_semana(actual)}", 8)
    columnas = [
        ("Posición", "0"), ("Producto", None), ("Categoría", None), ("Unidades", FORMATO_ENTERO),
        ("Ventas", FORMATO_PESOS), ("% del total", "0.0%"), ("Posición semana anterior", "0"), ("Cambio", None),
    ]
    registros = []
    colores = []
    for fila in ranking.itertuples():
        cambio, color = texto_cambio(fila.posicion, fila.posicion_anterior)
        registros.append([fila.posicion, fila.producto, fila.categoria, fila.unidades, fila.monto,
                          fila.porcentaje, fila.posicion_anterior, cambio])
        colores.append(color)
    fila_encabezado = 5
    escribir_tabla(hoja, fila_encabezado, columnas, registros)
    for i, color in enumerate(colores):
        celda = hoja.cell(fila_encabezado + 1 + i, 8)
        celda.font = Font(bold=True, color=color)
        celda.alignment = Alignment(horizontal="center")
    hoja.freeze_panes = hoja.cell(fila_encabezado + 1, 1)


def hoja_reposicion(hoja, tabla, sin_receta, actual):
    anchos = [18, 9, 12, 13, 11, 11, 18, 15]
    fijar_anchos(hoja, dict(zip("ABCDEFGH", anchos)))
    encabezado_hoja(hoja, f"Según el consumo de la semana del {rango_semana(actual)} y el stock al cierre del domingo", 8)
    columnas = [
        ("Insumo", None), ("Unidad", None), ("Stock al cierre", None), ("Consumo de la semana", None),
        ("Consumo diario", None), ("Días que alcanza", "0.0"), ("Estado", None), ("Cantidad sugerida a pedir", None),
    ]
    registros = []
    for fila in tabla.itertuples():
        registros.append([
            fila.insumo, fila.unidad, fila.stock, fila.consumo_semana, fila.consumo_diario,
            SIN_DATO if vacio(fila.dias_cobertura) else fila.dias_cobertura,
            fila.estado,
            SIN_DATO if vacio(fila.cantidad_sugerida) else fila.cantidad_sugerida,
        ])
    fila_encabezado = 5
    ultima = escribir_tabla(hoja, fila_encabezado, columnas, registros)

    for i, fila in enumerate(tabla.itertuples()):
        f = fila_encabezado + 1 + i
        formato = FORMATO_ENTERO if fila.unidad == "unidad" else FORMATO_DECIMAL
        for c in (3, 4, 5, 8):
            hoja.cell(f, c).number_format = formato
        for c in (6, 8):
            hoja.cell(f, c).alignment = Alignment(horizontal="right")
        fondo, texto = COLORES_ESTADO[fila.estado]
        estado = hoja.cell(f, 7)
        estado.fill = relleno(fondo)
        estado.font = Font(bold=True, color=texto)
        estado.alignment = Alignment(horizontal="center")
    hoja.freeze_panes = hoja.cell(fila_encabezado + 1, 1)

    notas = [
        f"Criterio: URGENTE si el stock alcanza para menos de {DIAS_URGENTE} días; Pedir si alcanza para menos "
        f"de {DIAS_PEDIR}; OK en otro caso.",
        f"Cantidad sugerida = consumo de {DIAS_A_CUBRIR + DIAS_MARGEN} días ({DIAS_A_CUBRIR} + {DIAS_MARGEN} "
        f"de margen por si el pedido se demora) − stock actual, redondeada hacia arriba. "
        f"Solo se sugiere para URGENTE y Pedir.",
    ]
    for producto in sin_receta.itertuples():
        notas.append(f"Atención: '{producto.producto}' ({producto.unidades} unidades vendidas) no descuenta "
                     f"insumos porque no tiene receta en el inventario.")
    escribir_notas(hoja, ultima + 2, notas, 8, sum(anchos))


def escribir_notas(hoja, fila, notas, ultima_columna, ancho_total):
    for i, nota in enumerate(notas):
        f = fila + i
        celda = hoja.cell(f, 1, nota)
        celda.font = Font(italic=True, color=GRIS_TEXTO)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        hoja.merge_cells(start_row=f, start_column=1, end_row=f, end_column=ultima_columna)
        hoja.row_dimensions[f].height = 15 * lineas_necesarias(nota, ancho_total)


def hoja_revision(hoja, registro):
    grupos = registro.agrupado()
    anchos = [26, 10, 32, 55, 42]
    fijar_anchos(hoja, dict(zip("ABCDE", anchos)))
    encabezado_hoja(
        hoja,
        f"{len(registro.correcciones)} correcciones agrupadas en {len(grupos)} líneas. "
        f"Ningún dato se cambió sin quedar registrado aquí.",
        5,
    )
    hoja["A4"] = ("'Filas' indica dónde está cada caso en el archivo original "
                  "(número de fila en Excel o de línea en el CSV).")
    hoja["A4"].font = Font(italic=True, color=GRIS_TEXTO)
    hoja.merge_cells("A4:E4")
    columnas = [("Archivo", None), ("N.º de filas", FORMATO_ENTERO), ("Filas", None), ("Problema", None), ("Acción", None)]
    registros = [[g["archivo"], g["cantidad"], g["filas"], g["problema"], g["accion"]] for g in grupos]
    escribir_tabla(hoja, 6, columnas, registros, ajustar_texto=True)
    hoja.freeze_panes = "A7"


def hoja_datos_limpios(hoja, ventas):
    anchos = [12, 14, 13, 20, 16, 10, 12, 27, 10]
    fijar_anchos(hoja, dict(zip("ABCDEFGHI", anchos)))
    encabezado_hoja(hoja, "Todas las ventas después de la limpieza, con su archivo y fila de origen", 9)
    columnas = [
        ("Fecha", FORMATO_FECHA), ("Semana (lunes)", FORMATO_FECHA), ("Sucursal", None), ("Producto", None),
        ("Categoría", None), ("Cantidad", FORMATO_ENTERO), ("Monto", FORMATO_PESOS),
        ("Archivo de origen", None), ("Fila de origen", "0"),
    ]
    registros = ventas[["fecha", "semana", "sucursal", "producto", "categoria",
                        "cantidad", "monto", "archivo", "fila_origen"]].values.tolist()
    fila_encabezado = 4
    ultima = escribir_tabla(hoja, fila_encabezado, columnas, registros)
    hoja.auto_filter.ref = f"A{fila_encabezado}:I{ultima}"
    hoja.freeze_panes = hoja.cell(fila_encabezado + 1, 1)


# --- Programa principal -------------------------------------------------------

def generar(carpeta=Path("datos_semana"), ruta_salida=Path("salida/reporte_semanal.xlsx")):
    """Genera el reporte y devuelve un resumen de lo que hizo."""
    carpeta, ruta_salida = Path(carpeta), Path(ruta_salida)
    inventario = carpeta / "inventario.xlsx"
    registro = Registro()
    ventas = cargar_ventas(carpeta, cargar_catalogo(inventario), registro)
    recetas, stock = cargar_inventario(inventario)
    actual, anterior = semanas(ventas)

    resumen = resumen_sucursales(ventas, actual, anterior)
    diarias = ventas_diarias(ventas, actual, anterior)
    ranking = ranking_productos(ventas, actual, anterior)
    tabla_reposicion = reposicion(ventas, recetas, stock, actual)
    sin_receta = productos_sin_receta(ventas, recetas, actual)

    libro = Workbook()
    libro.remove(libro.active)
    hojas = {nombre: libro.create_sheet(nombre) for nombre in HOJAS}
    for hoja in hojas.values():
        hoja.sheet_view.showGridLines = False
    hoja_resumen(hojas["Resumen"], resumen, diarias, ranking, actual, anterior)
    hoja_ranking(hojas["Ranking de productos"], ranking, actual, anterior)
    hoja_reposicion(hojas["Reposición de insumos"], tabla_reposicion, sin_receta, actual)
    hoja_revision(hojas["Revisión de datos"], registro)
    hoja_datos_limpios(hojas["Datos limpios"], ventas)

    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta_salida)

    total = resumen[resumen["sucursal"] == TOTAL].iloc[0]
    return {
        "ruta": ruta_salida,
        "archivos": ventas["archivo"].nunique(),
        "filas": len(ventas),
        "correcciones": len(registro.correcciones),
        "grupos": len(registro.agrupado()),
        "actual": actual,
        "anterior": anterior,
        "ventas": int(total["ventas"]),
        "variacion": total["variacion"],
        "estados": tabla_reposicion["estado"].value_counts().to_dict(),
    }


def main():
    parser = argparse.ArgumentParser(description="Genera el reporte semanal de ventas en Excel.")
    parser.add_argument("--datos", default="datos_semana", help="carpeta con las planillas de la semana")
    parser.add_argument("--salida", default="salida/reporte_semanal.xlsx", help="ruta del Excel a crear")
    argumentos = parser.parse_args()

    info = generar(argumentos.datos, argumentos.salida)
    estados = info["estados"]
    variacion_texto = "" if vacio(info["variacion"]) else f" ({porcentaje_texto(info['variacion'])} vs semana anterior)"
    print(f"{TITULO} (DATOS DE EJEMPLO)")
    print(f"Archivos de ventas leídos: {info['archivos']} (más inventario.xlsx)")
    print(f"Filas de ventas limpias:   {info['filas']}")
    print(f"Correcciones registradas:  {info['correcciones']} ({info['grupos']} líneas en 'Revisión de datos')")
    print(f"Semana del reporte:        {rango_semana(info['actual'])}")
    print(f"Ventas totales:            {pesos(info['ventas'])}{variacion_texto}")
    print(f"Insumos:                   {estados.get('URGENTE', 0)} URGENTE, {estados.get('Pedir', 0)} Pedir, "
          f"{estados.get('OK', 0)} OK")
    print(f"Reporte guardado en:       {info['ruta']}")


if __name__ == "__main__":
    main()
