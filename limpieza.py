"""
Lectura y limpieza de las planillas de ventas de cada sucursal.

Cada sucursal envía su planilla en un formato distinto. Este módulo las lee,
corrige los problemas conocidos y las une en una sola tabla. Toda corrección
queda anotada en un Registro (archivo, fila, problema, acción): nada se
cambia en silencio.
"""
import csv
import difflib
import io
import numbers
import re
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# Nombre estándar de cada columna y las formas en que aparece en las planillas
# (ya normalizadas: minúsculas, sin tildes ni símbolos).
SINONIMOS_COLUMNAS = {
    "fecha": ["fecha", "dia"],
    "producto": ["producto", "plato", "item"],
    "cantidad": ["cantidad", "cant", "unidades"],
    "precio": ["precio", "precio unitario", "precio unit"],
    "total": ["total", "monto", "total linea"],
}
COLUMNA_POR_SINONIMO = {s: col for col, sinonimos in SINONIMOS_COLUMNAS.items() for s in sinonimos}

NUMEROS_EN_PALABRAS = {
    "un": 1, "uno": 1, "una": 1, "dos": 2, "tres": 3, "cuatro": 4, "cinco": 5,
    "seis": 6, "siete": 7, "ocho": 8, "nueve": 9, "diez": 10, "once": 11,
    "doce": 12, "trece": 13, "catorce": 14, "quince": 15,
}

# Abreviaciones que NO se pueden deducir del nombre oficial (ej.: "bl" para
# "Barros Luco"). Se completan con el cliente. Las que sí se deducen
# ("Chorri", "Lomo p.", "Emp. pino") se resuelven solas por prefijos.
ABREVIACIONES = {}

UMBRAL_SIMILITUD = 0.75
FORMATOS_FECHA = [("%Y-%m-%d", None), ("%d/%m/%Y", "DD/MM/AAAA"), ("%d-%m-%Y", "DD-MM-AAAA")]
EXCEL_DIA_CERO = date(1899, 12, 30)

# Cómo se llama cada sucursal y cómo exporta sus ventas:
#   "producto_dia": una fila por producto y día (una fila repetida es un error).
#   "pedido": una fila por pedido (filas idénticas son ventas distintas).
SUCURSALES = {
    "providencia": ("Providencia", "producto_dia"),
    "nunoa": ("Ñuñoa", "producto_dia"),
    "maipu": ("Maipú", "pedido"),
}
COLUMNAS_LIMPIAS = ["fecha", "producto", "categoria", "cantidad", "monto", "fila_origen"]


# --- Registro de correcciones -------------------------------------------------

class Registro:
    """Junta las correcciones hechas a los datos: archivo, fila, problema, acción."""

    def __init__(self):
        self.correcciones = []

    def agregar(self, archivo, fila, problema, accion):
        self.correcciones.append({
            "archivo": archivo,
            "fila": int(fila),
            "problema": problema,
            "accion": accion,
        })

    def agrupado(self):
        """Une las correcciones iguales del mismo archivo en una sola línea."""
        grupos = {}
        for c in self.correcciones:
            grupos.setdefault((c["archivo"], c["problema"], c["accion"]), []).append(c["fila"])
        ordenados = sorted(grupos.items(), key=lambda g: (g[0][0], min(g[1])))
        return [
            {
                "archivo": archivo,
                "cantidad": len(filas),
                "filas": resumir_filas(filas),
                "problema": problema,
                "accion": accion,
            }
            for (archivo, problema, accion), filas in ordenados
        ]


def resumir_filas(filas, maximo=10):
    """[2, 3, 4, 9] -> '2-4, 9'. Si hay muchos tramos, muestra los primeros."""
    filas = sorted(set(filas))
    tramos = []
    inicio = anterior = filas[0]
    for fila in filas[1:] + [None]:
        if fila is not None and fila == anterior + 1:
            anterior = fila
            continue
        tramos.append(str(inicio) if inicio == anterior else f"{inicio}-{anterior}")
        if fila is not None:
            inicio = anterior = fila
    if len(tramos) > maximo:
        return ", ".join(tramos[:maximo]) + f" y {len(tramos) - maximo} tramos más"
    return ", ".join(tramos)


# --- Utilidades ---------------------------------------------------------------

def es_vacio(valor):
    if valor is None:
        return True
    if isinstance(valor, str):
        return not valor.strip()
    return isinstance(valor, float) and pd.isna(valor)


def normalizar_texto(texto):
    """'  Emp. PINO ' -> 'emp pino': minúsculas, sin tildes, sin símbolos."""
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texto).split())


def es_numero(valor):
    return isinstance(valor, numbers.Number) and not isinstance(valor, bool)


def pesos(monto):
    return "$" + f"{int(monto):,}".replace(",", ".")


# --- Lectura ------------------------------------------------------------------

def detectar_separador(texto, candidatos=(",", ";", "\t")):
    """Elige el separador que aparece la misma cantidad de veces en más líneas."""
    lineas = [l for l in texto.splitlines() if l.strip()][:20]
    mejor, mejor_puntaje = ",", (0, 0)
    for separador in candidatos:
        conteos = [l.count(separador) for l in lineas if l.count(separador) > 0]
        if not conteos:
            continue
        mas_comun = max(set(conteos), key=conteos.count)
        puntaje = (conteos.count(mas_comun), mas_comun)
        if puntaje > mejor_puntaje:
            mejor, mejor_puntaje = separador, puntaje
    return mejor


def normalizar_columnas(columnas):
    """Traduce cada nombre de columna a su nombre estándar; si no lo conoce, lo deja igual."""
    return [COLUMNA_POR_SINONIMO.get(normalizar_texto(c), c) for c in columnas]


def buscar_encabezado(filas, max_filas=15):
    """Índice de la primera fila con al menos 2 nombres de columna conocidos."""
    for i, fila in enumerate(filas[:max_filas]):
        conocidas = [c for c in fila if not es_vacio(c) and normalizar_texto(c) in COLUMNA_POR_SINONIMO]
        if len(conocidas) >= 2:
            return i
    raise ValueError(f"No se encontró la fila de encabezado en las primeras {max_filas} filas")


def leer_filas(ruta):
    """Todas las filas del archivo tal como vienen, sin interpretar nada."""
    ruta = Path(ruta)
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        hoja = load_workbook(ruta, data_only=True, read_only=True).active
        return [list(fila) for fila in hoja.iter_rows(min_row=1, values_only=True)]
    texto = ruta.read_text(encoding="utf-8-sig")
    return list(csv.reader(io.StringIO(texto), delimiter=detectar_separador(texto)))


def leer_planilla(ruta, registro):
    """Lee un CSV o Excel y devuelve una tabla con columnas estándar y la fila de origen.

    No corrige valores: solo ubica el encabezado y nombra las columnas.
    """
    ruta = Path(ruta)
    filas = leer_filas(ruta)
    i_encabezado = buscar_encabezado(filas)
    for i in range(i_encabezado):
        if any(not es_vacio(c) for c in filas[i]):
            registro.agregar(ruta.name, i + 1, "fila de título sobre el encabezado", "omitida")

    encabezado = filas[i_encabezado]
    indices = [j for j, c in enumerate(encabezado) if not es_vacio(c)]
    columnas = normalizar_columnas([encabezado[j] for j in indices])
    datos, numeros_fila = [], []
    for i in range(i_encabezado + 1, len(filas)):
        fila = filas[i]
        datos.append([fila[j] if j < len(fila) else None for j in indices])
        numeros_fila.append(i + 1)  # número de fila como se ve en Excel o en el editor

    tabla = pd.DataFrame(datos, columns=columnas, dtype=object)
    tabla["fila_origen"] = numeros_fila
    return tabla


# --- Valores ------------------------------------------------------------------

def interpretar_fecha(valor):
    """Devuelve (fecha, descripción del formato) o (None, None) si no se entiende.

    La descripción es None cuando el formato ya es el estándar y no hubo que corregir.
    """
    if isinstance(valor, datetime):
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None
    if es_numero(valor) and not pd.isna(valor):
        if 20000 <= valor <= 80000:
            return EXCEL_DIA_CERO + timedelta(days=int(valor)), "fecha guardada como número de Excel"
        return None, None
    if isinstance(valor, str):
        texto = valor.strip()
        for formato, nombre in FORMATOS_FECHA:
            try:
                fecha = datetime.strptime(texto, formato).date()
            except ValueError:
                continue
            return fecha, (f"fecha en formato {nombre}" if nombre else None)
    return None, None


def leer_fecha(valor):
    return interpretar_fecha(valor)[0]


def leer_numero(valor):
    """'$9.900' -> 9900, ' 12 ' -> 12, 'tres' -> 3. El punto es separador de miles."""
    if es_numero(valor):
        if pd.isna(valor):
            return None
        return int(valor) if float(valor).is_integer() else valor
    if not isinstance(valor, str):
        return None
    palabra = normalizar_texto(valor)
    if palabra in NUMEROS_EN_PALABRAS:
        return NUMEROS_EN_PALABRAS[palabra]
    texto = valor.replace("$", "").replace(" ", "").strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", texto):
        texto = texto.replace(".", "")
    if re.fullmatch(r"\d+", texto):
        return int(texto)
    return None


def es_abreviacion(palabras_abreviadas, palabras_producto):
    """Cada palabra abreviada es el inicio de una palabra del producto, en orden.

    La primera debe calzar con la primera: 'lomo p' -> 'lomo a lo pobre'.
    """
    if not palabras_abreviadas or not palabras_producto[0].startswith(palabras_abreviadas[0]):
        return False
    restantes = iter(palabras_producto[1:])
    return all(any(p.startswith(a) for p in restantes) for a in palabras_abreviadas[1:])


def emparejar_producto(nombre, catalogo):
    """Busca el producto oficial que corresponde a un nombre escrito a mano.

    Devuelve (producto, tipo). El tipo dice qué hubo que hacer:
    "exacto", "formato" (mayúsculas/tildes/espacios), "abreviación", "tipeo"
    o "desconocido" (producto None). Una abreviación que calza con más de un
    producto no se adivina.
    """
    if es_vacio(nombre):
        return None, "desconocido"
    texto = str(nombre)
    if texto in catalogo:
        return texto, "exacto"

    normalizado = normalizar_texto(texto)
    por_normalizado = {normalizar_texto(p): p for p in catalogo}
    if normalizado in por_normalizado:
        return por_normalizado[normalizado], "formato"

    if normalizado in ABREVIACIONES and ABREVIACIONES[normalizado] in catalogo:
        return ABREVIACIONES[normalizado], "abreviación"

    palabras = normalizado.split()
    candidatos = [p for n, p in por_normalizado.items() if es_abreviacion(palabras, n.split())]
    if len(candidatos) == 1:
        return candidatos[0], "abreviación"
    if len(candidatos) > 1:
        return None, "desconocido"

    parecidos = difflib.get_close_matches(normalizado, por_normalizado, n=1, cutoff=UMBRAL_SIMILITUD)
    if parecidos:
        return por_normalizado[parecidos[0]], "tipeo"
    return None, "desconocido"


# --- Limpieza de tablas -------------------------------------------------------

def quitar_filas_basura(tabla, archivo, registro):
    """Elimina filas vacías y filas de TOTAL escritas a mano (el total se recalcula)."""
    columnas = [c for c in tabla.columns if c != "fila_origen"]
    conservar = []
    for _, fila in tabla.iterrows():
        valores = [fila[c] for c in columnas]
        if all(es_vacio(v) for v in valores):
            registro.agregar(archivo, fila["fila_origen"], "fila vacía", "eliminada")
            conservar.append(False)
        elif any(isinstance(v, str) and normalizar_texto(v).split()[:1] == ["total"] for v in valores):
            registro.agregar(archivo, fila["fila_origen"], "fila de TOTAL escrita a mano",
                             "eliminada (el total se recalcula)")
            conservar.append(False)
        else:
            conservar.append(True)
    return tabla[conservar].reset_index(drop=True)


def leer_monto(fila, columna, archivo, registro):
    valor = fila.get(columna)
    monto = leer_numero(valor)
    if monto is not None and isinstance(valor, str) and ("$" in valor or "." in valor):
        registro.agregar(archivo, fila["fila_origen"], "monto con '$' y punto de miles",
                         "convertido a número")
    return monto


def limpiar_tabla(tabla, archivo, catalogo, registro):
    """Convierte cada fila a fecha, producto oficial, cantidad y monto, registrando cada cambio.

    Las filas sin fecha, cantidad, producto o monto legibles se descartan (y se registran).
    """
    limpias = []
    for fila in tabla.to_dict("records"):
        n = fila["fila_origen"]

        fecha, formato = interpretar_fecha(fila.get("fecha"))
        if fecha is None:
            registro.agregar(archivo, n, f"fecha ilegible: '{fila.get('fecha')}'", "fila descartada")
            continue
        if formato:
            registro.agregar(archivo, n, formato, "convertida a fecha")

        valor_cantidad = fila.get("cantidad")
        cantidad = leer_numero(valor_cantidad)
        if cantidad is None or cantidad <= 0:
            registro.agregar(archivo, n, f"cantidad ilegible: '{valor_cantidad}'", "fila descartada")
            continue
        if isinstance(valor_cantidad, str) and normalizar_texto(valor_cantidad) in NUMEROS_EN_PALABRAS:
            registro.agregar(archivo, n, f"cantidad escrita en palabras: '{valor_cantidad.strip()}'",
                             f"convertida a {cantidad}")

        escrito = fila.get("producto")
        if es_vacio(escrito):
            registro.agregar(archivo, n, "fila sin producto", "fila descartada")
            continue
        producto, tipo = emparejar_producto(escrito, catalogo)
        if tipo == "formato":
            registro.agregar(archivo, n, "nombre con mayúsculas, tildes o espacios distintos al catálogo",
                             "unificado con el nombre oficial")
        elif tipo == "tipeo":
            registro.agregar(archivo, n, f"nombre con tipeo: '{escrito}'", f"unificado como '{producto}'")
        elif tipo == "abreviación":
            registro.agregar(archivo, n, f"nombre abreviado: '{escrito}'",
                             f"abreviación expandida a '{producto}'")
        elif tipo == "desconocido":
            registro.agregar(archivo, n, f"producto desconocido: '{escrito}' (no está en el catálogo)",
                             "se conserva con su nombre, sin categoría")
            producto = str(escrito).strip()

        if "total" in fila:
            monto = leer_monto(fila, "total", archivo, registro)
        else:
            precio = leer_monto(fila, "precio", archivo, registro)
            monto = None if precio is None else cantidad * precio
        if monto is None:
            registro.agregar(archivo, n, "monto ilegible", "fila descartada")
            continue

        if producto in catalogo:
            esperado = cantidad * catalogo[producto]["precio"]
            if monto != esperado:
                registro.agregar(
                    archivo, n,
                    f"monto no calza con el catálogo: planilla {pesos(monto)}, catálogo {pesos(esperado)}",
                    "se respeta el monto de la planilla",
                )
            categoria = catalogo[producto]["categoria"]
        else:
            categoria = "Sin categoría"

        limpias.append({
            "fecha": fecha, "producto": producto, "categoria": categoria,
            "cantidad": cantidad, "monto": monto, "fila_origen": n,
        })
    return pd.DataFrame(limpias, columns=COLUMNAS_LIMPIAS)


def quitar_duplicados(tabla, formato, archivo, registro):
    """Quita filas repetidas solo cuando el formato es "una fila por producto y día".

    En el formato "pedido" dos filas iguales son dos ventas reales y se conservan.
    """
    if formato not in ("producto_dia", "pedido"):
        raise ValueError(f"Formato desconocido: '{formato}'. Usa 'producto_dia' o 'pedido'.")
    if formato == "pedido":
        return tabla.copy()

    clave = [c for c in tabla.columns if c != "fila_origen"]
    repetidas = tabla.duplicated(subset=clave, keep="first")
    primeras = {}
    for _, fila in tabla.iterrows():
        llave = tuple(fila[c] for c in clave)
        if llave not in primeras:
            primeras[llave] = fila["fila_origen"]
        else:
            registro.agregar(archivo, fila["fila_origen"],
                             f"fila duplicada (igual a la fila {primeras[llave]})", "eliminada")
    return tabla[~repetidas].reset_index(drop=True)


# --- Carga completa -----------------------------------------------------------

def cargar_catalogo(ruta_inventario):
    """{producto: {"categoria": ..., "precio": ...}} desde la hoja Catálogo."""
    hoja = pd.read_excel(ruta_inventario, sheet_name="Catálogo")
    return {
        fila["Producto"]: {"categoria": fila["Categoría"], "precio": int(fila["Precio"])}
        for _, fila in hoja.iterrows()
    }


def cargar_ventas(carpeta, catalogo, registro):
    """Lee, limpia y une las planillas de ventas de todas las sucursales.

    Devuelve una tabla con fecha, semana (lunes), sucursal, producto, categoría,
    cantidad, monto, archivo y fila de origen.
    """
    tablas = []
    for ruta in sorted(Path(carpeta).iterdir()):
        if ruta.suffix.lower() not in (".csv", ".xlsx") or ruta.name.startswith(("~$", ".")):
            continue
        if ruta.stem == "inventario":
            continue
        clave = ruta.stem.split("_")[0].lower()
        if clave not in SUCURSALES:
            registro.agregar(ruta.name, 0, "archivo de una sucursal desconocida", "ignorado completo")
            continue
        sucursal, formato = SUCURSALES[clave]

        tabla = leer_planilla(ruta, registro)
        tabla = quitar_filas_basura(tabla, ruta.name, registro)
        tabla = limpiar_tabla(tabla, ruta.name, catalogo, registro)
        tabla = quitar_duplicados(tabla, formato, ruta.name, registro)
        tabla["sucursal"] = sucursal
        tabla["archivo"] = ruta.name
        tablas.append(tabla)

    ventas = pd.concat(tablas, ignore_index=True)
    ventas["semana"] = [f - timedelta(days=f.weekday()) for f in ventas["fecha"]]
    return ventas[["fecha", "semana", "sucursal", "producto", "categoria",
                   "cantidad", "monto", "archivo", "fila_origen"]]
