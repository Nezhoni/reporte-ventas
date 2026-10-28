"""
Crea archivos de ejemplo desordenados para el reporte semanal de ventas.

Los datos son FICTICIOS. Simulan las planillas que envían cada semana las
3 sucursales de "Cocina del Valle" (Providencia, Ñuñoa y Maipú) y el
inventario central. Cada sucursal tiene su propio formato y sus propios
errores, para comprobar que la limpieza los detecta y los registra.

Uso:
    python3 crear_datos_ejemplo.py                   # crea datos_semana/
    python3 crear_datos_ejemplo.py --carpeta otra/   # crea los archivos en otra carpeta

Por seguridad, el script no escribe en una carpeta que ya tenga archivos:
los archivos de datos_semana/ son la entrada original y no se sobrescriben.
"""
import argparse
import csv
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

SEMILLA = 2026
# Lunes de cada semana: la primera es la "semana anterior", la segunda la "semana actual".
INICIO_SEMANAS = [date(2026, 10, 12), date(2026, 10, 19)]

# (producto, categoría, precio)
CATALOGO = [
    ("Chorrillana", "Platos de fondo", 12900),
    ("Lomo a lo pobre", "Platos de fondo", 11900),
    ("Lomo saltado", "Platos de fondo", 10900),
    ("Pastel de choclo", "Platos de fondo", 9900),
    ("Cazuela de vacuno", "Platos de fondo", 8900),
    ("Empanada de pino", "Entradas", 2900),
    ("Empanada de queso", "Entradas", 2500),
    ("Sopaipillas", "Entradas", 3500),
    ("Completo italiano", "Sándwiches", 4500),
    ("Barros Luco", "Sándwiches", 7900),
    ("Bebida en lata", "Bebidas", 1900),
    ("Jugo natural", "Bebidas", 2900),
    ("Pisco sour", "Bebidas", 4900),
]
PRECIOS = {producto: precio for producto, _, precio in CATALOGO}
PRECIOS_MAYUSCULAS = {producto.upper(): precio for producto, precio in PRECIOS.items()}

# Unidades vendidas por día en una sucursal "promedio".
POPULARIDAD = {
    "Chorrillana": 9, "Lomo a lo pobre": 7, "Lomo saltado": 6,
    "Pastel de choclo": 5, "Cazuela de vacuno": 5, "Empanada de pino": 14,
    "Empanada de queso": 8, "Sopaipillas": 6, "Completo italiano": 7,
    "Barros Luco": 5, "Bebida en lata": 18, "Jugo natural": 9, "Pisco sour": 8,
}
FACTOR_SUCURSAL = {"Providencia": 1.2, "Ñuñoa": 1.0, "Maipú": 0.85}
CRECIMIENTO_SEMANA_ACTUAL = {"Providencia": 1.08, "Ñuñoa": 0.95, "Maipú": 1.03}
FACTOR_DIA = [1.0, 0.9, 1.0, 1.1, 1.4, 1.5, 1.2]  # lunes a domingo

# Insumo y su unidad de medida.
UNIDADES = {
    "Papas": "kg", "Carne de vacuno": "kg", "Cebolla": "kg", "Huevos": "unidad",
    "Choclo": "kg", "Pollo": "kg", "Arroz": "kg", "Tomate": "kg",
    "Harina": "kg", "Queso": "kg", "Aceite": "litro", "Pan": "unidad",
    "Vienesas": "unidad", "Palta": "kg", "Bebidas en lata": "unidad",
    "Fruta": "kg", "Pisco": "litro", "Limón": "kg",
}

# Cantidad de cada insumo que lleva UNA unidad vendida del producto.
RECETAS = {
    "Chorrillana": {"Papas": 0.35, "Carne de vacuno": 0.2, "Cebolla": 0.1, "Huevos": 2},
    "Lomo a lo pobre": {"Carne de vacuno": 0.25, "Papas": 0.25, "Huevos": 2, "Cebolla": 0.08},
    "Lomo saltado": {"Carne de vacuno": 0.2, "Papas": 0.15, "Cebolla": 0.1, "Tomate": 0.1, "Arroz": 0.12},
    "Pastel de choclo": {"Choclo": 0.3, "Carne de vacuno": 0.1, "Pollo": 0.08, "Cebolla": 0.05, "Huevos": 0.5},
    "Cazuela de vacuno": {"Carne de vacuno": 0.2, "Papas": 0.15, "Choclo": 0.1, "Arroz": 0.03},
    "Empanada de pino": {"Harina": 0.12, "Carne de vacuno": 0.08, "Cebolla": 0.06, "Huevos": 0.25},
    "Empanada de queso": {"Harina": 0.12, "Queso": 0.08},
    "Sopaipillas": {"Harina": 0.15, "Aceite": 0.05},
    "Completo italiano": {"Pan": 1, "Vienesas": 1, "Palta": 0.08, "Tomate": 0.05},
    "Barros Luco": {"Pan": 1, "Carne de vacuno": 0.15, "Queso": 0.06},
    "Bebida en lata": {"Bebidas en lata": 1},
    "Jugo natural": {"Fruta": 0.3},
    "Pisco sour": {"Pisco": 0.06, "Limón": 0.05, "Huevos": 0.2},
}

# Días que alcanza el stock de cada insumo con el consumo de la semana actual.
# Calibrado para que haya pocas alertas: 2 urgentes (< 2 días), 3 por pedir (< 5 días).
DIAS_COBERTURA = {
    "Palta": 1.3, "Limón": 1.6,                      # URGENTE
    "Carne de vacuno": 3.4, "Huevos": 2.7, "Pan": 4.1,  # Pedir
    "Papas": 6.5, "Cebolla": 9, "Choclo": 7.5, "Pollo": 10, "Arroz": 14,
    "Tomate": 6, "Harina": 8, "Queso": 11, "Aceite": 12, "Vienesas": 7,
    "Bebidas en lata": 9, "Fruta": 6.5, "Pisco": 13,
}

# Cómo anota los platos el personal de Ñuñoa: (nombre escrito, probabilidad).
# Lo que no suma 1 se escribe con el nombre oficial.
VARIANTES_NUNOA = {
    "Chorrillana": [("Chorillana", 0.2), ("Chorri", 0.2)],
    "Lomo a lo pobre": [("Lomo p.", 0.3)],
    "Lomo saltado": [("lomo saltao", 0.25)],
    "Empanada de pino": [("Emp. pino", 0.3)],
}
NUMEROS_EN_PALABRAS = {
    1: "uno", 2: "dos", 3: "tres", 4: "cuatro", 5: "cinco",
    6: "seis", 7: "siete", 8: "ocho", 9: "nueve", 10: "diez",
}
PRODUCTO_FUERA_DE_CATALOGO = ("Promo almuerzo", 6900)


def generar_ventas_reales(semilla=SEMILLA):
    """Ventas "verdaderas", sin ningún error: una por producto, día y sucursal.

    Son la base de los archivos desordenados. También sirven para comprobar
    después que la limpieza recupera exactamente estas cifras.
    """
    azar = random.Random(semilla)
    ventas = []
    for n_semana, inicio in enumerate(INICIO_SEMANAS):
        for sucursal, factor in FACTOR_SUCURSAL.items():
            crecimiento = CRECIMIENTO_SEMANA_ACTUAL[sucursal] if n_semana == 1 else 1.0
            for d in range(7):
                for producto, _, _ in CATALOGO:
                    media = POPULARIDAD[producto] * factor * crecimiento * FACTOR_DIA[d]
                    cantidad = max(0, round(azar.gauss(media, media * 0.25)))
                    if cantidad:
                        ventas.append({
                            "fecha": inicio + timedelta(days=d),
                            "sucursal": sucursal,
                            "producto": producto,
                            "cantidad": cantidad,
                        })
    return ventas


def calcular_stock(ventas):
    """Stock al cierre del domingo, calibrado con DIAS_COBERTURA."""
    inicio_actual = INICIO_SEMANAS[-1]
    consumo = {insumo: 0.0 for insumo in UNIDADES}
    for venta in ventas:
        if venta["fecha"] >= inicio_actual:
            for insumo, por_unidad in RECETAS[venta["producto"]].items():
                consumo[insumo] += venta["cantidad"] * por_unidad
    filas = []
    for insumo, unidad in UNIDADES.items():
        stock = consumo[insumo] / 7 * DIAS_COBERTURA[insumo]
        stock = round(stock) if unidad == "unidad" else round(stock, 1)
        filas.append((insumo, unidad, stock))
    return filas


def pesos(monto):
    """12900 -> '$12.900' (formato chileno, punto de miles)."""
    return "$" + f"{monto:,}".replace(",", ".")


def fecha_excel(dia):
    """Número de serie con que Excel guarda las fechas (días desde 30/12/1899)."""
    return (dia - date(1899, 12, 30)).days


def ventas_de(ventas, sucursal, inicio):
    fin = inicio + timedelta(days=7)
    return [v for v in ventas if v["sucursal"] == sucursal and inicio <= v["fecha"] < fin]


def escribir_providencia(ventas, inicio, ruta, azar):
    """CSV con ';', títulos arriba, montos '$9.900', MAYÚSCULAS y fila TOTAL a mano."""
    lineas = []
    for v in ventas_de(ventas, "Providencia", inicio):
        precio = PRECIOS[v["producto"]]
        lineas.append([
            v["fecha"].strftime("%d/%m/%Y"),
            v["producto"].upper(),
            str(v["cantidad"]),
            pesos(precio),
            pesos(precio * v["cantidad"]),
        ])
    desordenes = [f"títulos sobre el encabezado, montos con '$' y punto de miles, "
                  f"{len(lineas)} nombres en MAYÚSCULAS, fila TOTAL a mano"]
    if inicio == INICIO_SEMANAS[-1]:
        i = azar.randrange(len(lineas))
        lineas.insert(i + 1, list(lineas[i]))
        desordenes.append(f"fila duplicada por error: {lineas[i][0]} {lineas[i][1]}")
    # Total sumado a mano, incluida la fila repetida (como lo haría una persona).
    total = sum(PRECIOS_MAYUSCULAS[l[1]] * int(l[2]) for l in lineas)
    fin = inicio + timedelta(days=6)
    with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.writer(archivo, delimiter=";")
        escritor.writerow(["Cocina del Valle - Sucursal Providencia"])
        escritor.writerow([f"Ventas del {inicio:%d/%m/%Y} al {fin:%d/%m/%Y}"])
        escritor.writerow(["Fecha", "Plato", "Cantidad", "Precio unitario", "Total"])
        escritor.writerows(lineas)
        escritor.writerow(["", "TOTAL", "", "", pesos(total)])
    return desordenes


def elegir_nombre_nunoa(producto, azar):
    sorteo = azar.random()
    acumulado = 0.0
    for variante, probabilidad in VARIANTES_NUNOA.get(producto, []):
        acumulado += probabilidad
        if sorteo < acumulado:
            return variante
    return producto


def escribir_nunoa(ventas, inicio, ruta, azar):
    """Excel con otros nombres de columna, fechas como número, tipeos,
    abreviaciones, cantidades en palabras y filas vacías."""
    filas = []
    for v in ventas_de(ventas, "Ñuñoa", inicio):
        filas.append([
            fecha_excel(v["fecha"]),
            elegir_nombre_nunoa(v["producto"], azar),
            v["cantidad"],
            PRECIOS[v["producto"]] * v["cantidad"],
        ])
    # 4 cantidades escritas en palabras.
    candidatas = [i for i, fila in enumerate(filas) if fila[2] <= 10]
    for i in azar.sample(candidatas, 4):
        filas[i][2] = NUMEROS_EN_PALABRAS[filas[i][2]]
    # 2 filas vacías en medio de la planilla.
    for i in sorted(azar.sample(range(1, len(filas)), 2), reverse=True):
        filas.insert(i, [])

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas"
    hoja.append(["Día", "Item", "Cant.", "Total"])
    for fila in filas:
        hoja.append(fila)
    # Fechas fijas en las propiedades del archivo para que salga igual cada vez.
    libro.properties.created = datetime(2026, 10, 26, 9, 0)
    libro.properties.modified = datetime(2026, 10, 26, 9, 0)
    libro.save(ruta)

    nombres = [f[1] for f in filas if f]
    tipeos = sum(nombres.count(n) for n in ("Chorillana", "lomo saltao"))
    abreviados = sum(nombres.count(n) for n in ("Chorri", "Lomo p.", "Emp. pino"))
    return [
        "columnas 'Día', 'Item', 'Cant.', fechas como número de Excel",
        f"{tipeos} nombres con tipeo, {abreviados} platos abreviados",
        "4 cantidades en palabras, 2 filas vacías",
    ]


def escribir_maipu(ventas, inicio, ruta, azar):
    """CSV con tabulador, una fila por pedido (filas idénticas legítimas),
    fechas mezcladas y un producto que no está en el catálogo."""
    filas = []
    for d in range(7):
        dia = inicio + timedelta(days=d)
        del_dia = []
        for v in ventas_de(ventas, "Maipú", inicio):
            if v["fecha"] != dia:
                continue
            resto = v["cantidad"]
            while resto > 0:
                cantidad = 2 if resto >= 2 and azar.random() < 0.3 else 1
                del_dia.append([dia, v["producto"], cantidad, PRECIOS[v["producto"]]])
                resto -= cantidad
        if d < 5:  # la promo solo existe de lunes a viernes
            nombre, precio = PRODUCTO_FUERA_DE_CATALOGO
            del_dia.append([dia, nombre, 1, precio])
        azar.shuffle(del_dia)
        filas.extend(del_dia)

    n_otro_formato = 0
    with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.writer(archivo, delimiter="\t")
        escritor.writerow(["fecha", "producto", "cantidad", "precio"])
        for dia, producto, cantidad, precio in filas:
            if azar.random() < 0.12:
                texto_fecha = dia.strftime("%d-%m-%Y")
                n_otro_formato += 1
            else:
                texto_fecha = dia.strftime("%Y-%m-%d")
            escritor.writerow([texto_fecha, producto, cantidad, precio])

    claves = [tuple(f) for f in filas]
    repetidas = len(claves) - len(set(claves))
    return [
        f"{n_otro_formato} fechas en formato DD-MM-AAAA (el resto AAAA-MM-DD)",
        f"{repetidas} filas idénticas legítimas (pedidos distintos, NO se borran)",
        f"5 filas de '{PRODUCTO_FUERA_DE_CATALOGO[0]}', que no está en el catálogo",
    ]


def escribir_inventario(ventas, ruta):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Catálogo"
    hoja.append(["Producto", "Categoría", "Precio"])
    for fila in CATALOGO:
        hoja.append(list(fila))

    hoja = libro.create_sheet("Recetas")
    hoja.append(["Producto", "Insumo", "Cantidad por unidad"])
    for producto, insumos in RECETAS.items():
        for insumo, cantidad in insumos.items():
            hoja.append([producto, insumo, cantidad])

    hoja = libro.create_sheet("Stock")
    hoja.append(["Insumo", "Unidad", "Stock al cierre del domingo"])
    for fila in calcular_stock(ventas):
        hoja.append(list(fila))

    libro.properties.created = datetime(2026, 10, 26, 9, 0)
    libro.properties.modified = datetime(2026, 10, 26, 9, 0)
    libro.save(ruta)


def crear_datos(carpeta):
    ventas = generar_ventas_reales()
    resumen = []
    for inicio in INICIO_SEMANAS:
        sufijo = inicio.isoformat()
        archivos = [
            (f"providencia_{sufijo}.csv", escribir_providencia),
            (f"nunoa_{sufijo}.xlsx", escribir_nunoa),
            (f"maipu_{sufijo}.csv", escribir_maipu),
        ]
        for nombre, escribir in archivos:
            # Un generador aleatorio por archivo: cambiar uno no altera los demás.
            azar = random.Random(f"{SEMILLA}-{nombre}")
            resumen.append((nombre, escribir(ventas, inicio, carpeta / nombre, azar)))
    escribir_inventario(ventas, carpeta / "inventario.xlsx")
    resumen.append(("inventario.xlsx", ["hojas Catálogo, Recetas y Stock"]))
    return resumen


def main():
    parser = argparse.ArgumentParser(description="Crea datos de ejemplo (ficticios).")
    parser.add_argument("--carpeta", default="datos_semana", help="carpeta de destino")
    carpeta = Path(parser.parse_args().carpeta)

    if carpeta.exists() and any(carpeta.iterdir()):
        sys.exit(f"La carpeta '{carpeta}' ya tiene archivos. No se sobrescriben "
                 f"porque son la entrada original. Bórrala a mano si quieres regenerarla.")
    carpeta.mkdir(parents=True, exist_ok=True)

    print(f"Creando datos de ejemplo (FICTICIOS) en {carpeta}/\n")
    for nombre, desordenes in crear_datos(carpeta):
        print(nombre)
        for desorden in desordenes:
            print(f"  - {desorden}")


if __name__ == "__main__":
    main()
